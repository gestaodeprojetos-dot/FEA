"""Embute os vídeos de animação no pitch do designer, um por slide, tocando sozinhos.

Uso: python3 fea_montar_pptx.py ENTRADA.pptx PASTA_ANIM PASTA_IMAGENS SAIDA.pptx
Cada vídeo cobre o slide inteiro, começa automaticamente ao entrar no slide,
não fica oculto e para no último quadro (que é a própria arte do designer).
"""
import glob
import re
import subprocess
import sys

from lxml import etree
from pptx import Presentation

SRC, ANIM, IMG, OUT = sys.argv[1:5]
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'

TIMING = '''<p:timing xmlns:p="%s"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="%d" fill="hold"/><p:tgtEl><p:spTgt spid="%s"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="0" mute="1"><p:cTn id="7" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst></p:cTn><p:tgtEl><p:spTgt spid="%s"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''

prs = Presentation(SRC)
W, H = prs.slide_width, prs.slide_height
done = []
for f in sorted(glob.glob(f'{ANIM}/FEA-elite-anim-slide-*.mp4')):
    n = int(re.search(r'slide-(\d+)', f).group(1))
    slide = prs.slides[n - 1]
    dur = int(float(subprocess.check_output(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'csv=p=0', f])) * 1000)
    # capa do vídeo = slide anterior, que é o que já está na tela quando o vídeo começa
    mv = slide.shapes.add_movie(f, 0, 0, W, H, poster_frame_image=f'{IMG}/s{n - 1:03d}.jpg', mime_type='video/mp4')
    sid = mv.shape_id
    old = slide._element.find(f'{{{P}}}timing')
    if old is not None:
        slide._element.remove(old)
    new = etree.fromstring(TIMING % (P, dur, sid, sid))
    # p:timing precisa vir depois de p:clrMapOvr / p:transition, antes de p:extLst
    ext = slide._element.find(f'{{{P}}}extLst')
    if ext is not None:
        ext.addprevious(new)
    else:
        slide._element.append(new)
    done.append(n)
prs.save(OUT)
print(len(done), 'slides com animação:', done)
