#!/usr/bin/env python3
"""Planilha "Links úteis" LATAM no padrão exato da planilha Brasil (Links úteis | Perpétuo).

Usa o próprio xlsx exportado do Brasil como molde (fonte, larguras, estilos, cabeçalho)
e troca só o conteúdo da coluna LINK: texto = nome do arquivo, hyperlink = URL.
    python3 fea_links_uteis_modelo.py br.xlsx SAIDA.xlsx
"""
import sys, copy
from openpyxl import load_workbook
from openpyxl.styles import Font

D = 'https://drive.google.com/drive/folders/'
G = 'https://docs.google.com/document/d/'
S = 'https://docs.google.com/spreadsheets/d/'
F = 'https://drive.google.com/file/d/'

# linha do molde Brasil -> (texto exibido, url ou None)
LINHAS = {
    2: ('Ebook olheiras Latam', D + '1n2cSn7QoPj8rxcp7MvZTNAuZOzVpY4kc'),
    4: ('1. COPY', D + '1YUTK9DaHqRcbeJq79ZWNHMA6StfcghCr'),
    5: ('FEA-ES | Copy Página eBook Ojeras', G + '1v-gq1sJjdatMomxQkThTJSERjOuFqiUAlBjroz7u0no/edit'),
    6: ('FEA-ES | Página de Upsell - Cpx', G + '1aSJo7XbjbAOpMuFhbrYG1a_5BKd6y5jxzyvhUjdApgg/edit'),
    7: ('FEA-ES | Checkout personalizado I e-Book', G + '13TzqiYe3ADx0CGK7hBffPOMh1W_HV3LW18RsTmSjcfQ/edit'),
    8: ('FEA-ES | Creativos - eBook Relleno de Ojeras', G + '1_QbWisDRpgBmrsWcSjPOU45LBBQ0Ce8UpUjViRiypGg/edit'),
    9: ('FEA-ES | Creativos - eBook Relleno de Ojeras 2', G + '1gJ1mNDlW2-jjv8WSXxI64cES_M_enf52hoL3C4zdsyg/edit'),
    10: ('FEA-ES | Leyenda I Creativos del eBook Ojeras', G + '1UgSAaYLfayStlF23n3jkKLxqM3u3I4481OtwTw4qPGg/edit'),
    11: ('FEA-ES | Listboss API - eBook Ojeras', G + '1a1HERGQ99ZMg6Yv_Ku6SUfI4kLi0HPrXYV3FD1cxMGQ/edit'),
    12: ('FEA-ES | Listboss Email - eBook Ojeras', G + '14_pPw0qV88mc0Z4dX1v_SWjx9kDpUrYImIe9908fBX4/edit'),
    13: ('FEA-ES | Descripción del grupo I E-book Ojeras', G + '11ScQrS6ZEnHQmYFlk_bx0XXIxlznbIUYqOln1C9_Sfk/edit'),
    14: ('FEA-Avatar Grupo de Alumnos - LATAM.png (pasta 3. Capas grupo e formulários)', D + '1ar8bBDc_TqFNULr1eawmTSgkfVFDQUrH'),
    16: ('2. ARTES', D + '1Q4UHzgdjlUx8_ewd1sBNWemyJ9JZT1ui'),
    17: ('1. Banner checkout', D + '1q5eyS55FaIDXOrsULKTYK2dAhL_jPlMq'),
    18: ('ESTÁTICOS', D + '1RibhoTeWsYU5wf4q__1aV_QGXGCia2n4'),
    19: ('VÍDEOS', D + '1ZKGD-pQ0iA8FHqx6Ekc8HvIWYeqkfPS3'),
    22: ('pendente: página de vendas LATAM (André Vivas)', None),
    23: ('pendente: página de upsell LATAM (André Vivas)', None),
    24: ('pendente: checkout Hotmart em USD', None),
    25: ('pendente: área de membros LATAM', None),
    29: ('pendente: páginas de teste LATAM', None), 30: (None, None), 31: (None, None), 32: (None, None),
    34: ('pendente: planilha de respostas da pesquisa LATAM', None),
    36: ('pendente: grupo de alunos LATAM', None),
}
EXTRAS = [  # mesmas duas colunas, depois do bloco do modelo
    ('Outros LATAM', None, None),
    ('Logo', 'FEA-Logo - Nome - Fonte - Ojeras LATAM.png (pasta 5. LOGO)', D + '1SNsyj8xq-LdMXNkmI53NdVPOQmwwjupi'),
    ('Ebook em espanhol (PDF)', 'Relleno_Tridimensional_de_Ojeras_ES.pdf', F + '1xZR8Y-3gtfZX9pnFUskn27V0O4kUw9al/view'),
    ('Aulas traduzidas', 'Aulas traduzidas', D + '1lJlpRTaLmGc0PGEau01jTusnW8uV5t91'),
    ('Copy dos criativos estáticos', 'FEA-LATAM Copy dos Criativos Estáticos (PT | ES) v2 - Ebook Olheiras', G + '159seUx0p5uXMYmNENALTJBuB7QV1pv3KKey2U_wDD4c/edit'),
    ('Planilha de tráfego', 'FEA-Criativos para Tráfego - Ebook Olheiras LATAM', S + '1ZxfcVB5kbLx_qagtWL0bk8OlBoqiDh2CnIVAGLA1PXo/edit'),
    ('Links dos vídeos traduzidos', 'FEA-Links dos vídeos traduzidos - Ebook Olheiras LATAM', S + '1k9Cg8s6SIRdDhgVgFWhp9aPUL04zEfRvUGMPQ0nU3kc/edit'),
    ('Páginas em HTML e JSON (designer)', 'Copys páginas ES', D + '15RFlmEyX_sWqlCHJuL8ZPEo-YA0m8TgB'),
    ('Briefing reverso', 'FEA-Briefing Reverso - Ebook Olheiras LATAM (preenchido)', G + '1EfDIyPQqWMK41_sL5AHsEBl6W5bXzXAUnE2ALu2nnO0/edit'),
    ('Backlog ClickUp', 'Ebook Olheiras LATAM', 'https://app.clickup.com/9013080622/v/f/901319449714/901314001225'),
]

wb = load_workbook(sys.argv[1])
for n in wb.sheetnames:
    if n != 'EBOOK OLHEIRAS':
        del wb[n]
ws = wb['EBOOK OLHEIRAS']
ws.title = 'EBOOK OLHEIRAS LATAM'
link_font = ws['B5'].font
pend_font = Font(name='Arial', color='FFCC0000')
for r, (txt, url) in LINHAS.items():
    c = ws.cell(r, 2)
    c.hyperlink = None
    c.value = txt
    if url:
        c.hyperlink = url
        c.font = copy.copy(link_font)
    elif txt:
        c.font = pend_font
base = 39
cab_font = ws['A4'].font
for i, (nome, txt, url) in enumerate(EXTRAS):
    r = base + i
    a, b = ws.cell(r, 1), ws.cell(r, 2)
    a.value = nome
    a.font = copy.copy(cab_font if url is None else ws["A5"].font)
    a.alignment = copy.copy(ws["A5"].alignment)
    if txt:
        b.value = txt
        b.hyperlink = url
        b.font = copy.copy(link_font)
        b.alignment = copy.copy(ws["B5"].alignment)
wb.save(sys.argv[2])
print('ok', sys.argv[2])
