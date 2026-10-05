#!/usr/bin/env python3
"""Gera as planilhas padrão de um novo evento FEA a partir de um JSON.

    python3 fea_planilhas_evento.py config.json SAIDA_DIR

Saída (xlsx, para subir no Drive com conversão para Google Sheets):
    FEA-Links úteis - <evento>.xlsx
    FEA-Criativos para Tráfego - <evento>.xlsx

Formato do config.json:
{
  "evento": "Ebook Olheiras LATAM",
  "links": [ {"secao": "GERAL", "linhas": [["Nome", "url ou texto", "url referência", "status"], ...]}, ... ],
  "trafego": {
    "campanha": "Ebook Olheiras LATAM",
    "meio": "Meta Ads",
    "link_cta": "pendente",
    "grupos": [ {"fase": "Vendas", "prefixo": "Ads", "sufixo": "PTO-LATAM", "qtd": 10,
                 "legenda": "url", "pasta": "url"} ]
  }
}
Link vira fórmula HYPERLINK com rótulo curto; texto sem http fica como texto.
Sem datas: a coluna Deadline sai vazia de propósito (preenchida no sprint).
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROXO, LAVANDA, CINZA, VERMELHO = "5B4599", "D0D1E9", "E9E9E9", "BB3838"
borda = Border(*(Side(style="thin", color=CINZA),) * 4)


def cabecalho(ws, colunas, larguras):
    ws.append(colunas)
    for i, c in enumerate(ws[ws.max_row], 1):
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ROXO)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[c.column_letter].width = larguras[i - 1]
    ws.freeze_panes = "A2"


def celula_link(c, valor, rotulo):
    if isinstance(valor, str) and valor.startswith("http"):
        c.value = f'=HYPERLINK("{valor}","{rotulo.replace(chr(34), chr(39))}")'
        c.font = Font(color="3846BB", underline="single")
    else:
        c.value = valor
        if isinstance(valor, str) and valor.lower().startswith("pendente"):
            c.font = Font(color=VERMELHO, bold=True)


def links(cfg, destino):
    wb = Workbook()
    ws = wb.active
    ws.title = "LINKS ÚTEIS"
    cabecalho(ws, ["NOME", "LINK", "REFERÊNCIA (evento modelo)", "STATUS"], [42, 55, 45, 30])
    for secao in cfg["links"]:
        ws.append([secao["secao"]])
        for c in ws[ws.max_row][:4]:
            c.font = Font(bold=True, color=ROXO)
            c.fill = PatternFill("solid", fgColor=LAVANDA)
        for nome, link, ref, status in secao["linhas"]:
            ws.append([nome])
            linha = ws[ws.max_row]
            celula_link(linha[1], link, f"Abrir · {nome}")
            celula_link(linha[2], ref, "Ver modelo")
            linha[3].value = status
            if status and "pendente" in status.lower():
                linha[3].font = Font(color=VERMELHO, bold=True)
        ws.append([])
    for row in ws.iter_rows(min_row=1, max_col=4):
        for c in row:
            c.border = borda
            c.alignment = Alignment(vertical="center", wrap_text=True)
    wb.save(os.path.join(destino, f"FEA-Links úteis - {cfg['evento']}.xlsx"))


def trafego(cfg, destino):
    t = cfg["trafego"]
    wb = Workbook()
    ws = wb.active
    ws.title = "CRIATIVOS"
    cols = ["Status", "Campanha", "Meio", "Deadline", "Fase", "Nome do criativo",
            "Legenda", "Link Criativo", "Link CTA", "UTM", "Observações"]
    cabecalho(ws, cols, [16, 30, 12, 12, 16, 32, 22, 22, 30, 22, 30])
    for g in t["grupos"]:
        for n in range(1, g["qtd"] + 1):
            for formato in ("Feed", "Story"):
                ws.append(["A produzir", t["campanha"], t["meio"], None, g["fase"],
                           f"{g['prefixo']} {n:02d} {formato} - {g['sufixo']}"])
                linha = ws[ws.max_row]
                celula_link(linha[6], g["legenda"], "Doc de legendas")
                celula_link(linha[7], g["pasta"], "Pasta do criativo")
                celula_link(linha[8], t["link_cta"], "Página")
                linha[9].value = "pendente (Matheus Saar)"
                linha[9].font = Font(color=VERMELHO)
    for row in ws.iter_rows(min_row=2, max_col=len(cols)):
        for c in row:
            c.border = borda
    leg = wb.create_sheet("COMO PREENCHER")
    cabecalho(leg, ["Status", "Quando usar"], [20, 80])
    for s in [("A produzir", "copy ou arte ainda não entregue"),
              ("Em revisão", "criativo pronto aguardando aprovação"),
              ("Agendar", "aprovado, falta subir no gerenciador"),
              ("No ar", "anúncio ativo"),
              ("Pausado", "desligado por performance ou fim de fase")]:
        leg.append(s)
    leg.append([])
    leg.append(["Regra", "Deadline fica vazio até a gestão definir o sprint. Nome do criativo igual ao arquivo salvo na pasta."])
    wb.save(os.path.join(destino, f"FEA-Criativos para Tráfego - {cfg['evento']}.xlsx"))


if __name__ == "__main__":
    cfg = json.load(open(sys.argv[1], encoding="utf-8"))
    os.makedirs(sys.argv[2], exist_ok=True)
    links(cfg, sys.argv[2])
    trafego(cfg, sys.argv[2])
    print("ok:", os.listdir(sys.argv[2]))
