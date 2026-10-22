"""Ajustes finais nas copys de disparo do MTB 11/26 (depois da troca de datas):
redistribui os disparos do primeiro dia (02/11), corrige cabeçalhos que a troca automática
não alcançou, tira o número de inscritos do 08/26 e aplica a regra de escrita FEA
("para" por extenso, sem travessão). Edita trecho a trecho, preservando a formatação.

Uso: python3 fea_ajustar_disparos_mtb_1126.py [--aplicar]"""
import io, re, sys
import docx
sys.path.insert(0, "/home/user/FEA/FEA-edicao-videos")
from fea_drive import token, chamar, API, UPLOAD

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
GDOC = "application/vnd.google-apps.document"
# id no Drive -> (nome, mime, trocas exatas)
DOCS = {
    "1K8oZeKQC6TiFdb9oHnzc26kzafD2dIlSpzrbPhu4QJ4": ("Captação WPP", GDOC, [
        ("02. Captação, segunda-feira, 02, 20h00", "02. Captação, segunda-feira, 02, 15h00"),
        ("03. Captação, segunda-feira, 02, 11h30", "03. Captação, segunda-feira, 02, 18h00"),
        ("04. Captação, segunda-feira, 02,  20h", "04. Captação, segunda-feira, 02, 20h30"),
        ("v 17. Captação", "💻  17. Captação"),
        ("As 9 técnicas avançadas — estéticas e funcionais", "As 9 técnicas avançadas, estéticas e funcionais"),
    ]),
    "1PGhkfRAZ95dzhlaoFOfMr_k5tn2l6UrC": ("E-mail Captação", DOCX, [
        ("02. Email, Captação, segunda-feira, 02, 07h", "02. Email, Captação, segunda-feira, 02, 19h"),
        ("Masterclass em Toxina Botulínica — um encontro", "Masterclass em Toxina Botulínica: um encontro"),
    ]),
    "1pwZYFIql4e4q87LCPq6mnJuqrm2dgUPS": ("E-mail Inscrições abertas FEB", DOCX, [
        ("01. Email, Inscrições abertas, Terça-feira, DISPARO MANUAL", "01. Email, Inscrições abertas, Segunda-feira, 09, DISPARO MANUAL"),
        ("Anatomia em Cadáveres – válido", "Anatomia em Cadáveres: válido"),
    ]),
    "1aKDkLWVN4RpbalALIJCZ_8I0xrJjgA7m": ("WhatsApp Inscrições abertas FEB", DOCX, [
        ("Instagram – válido", "Instagram: válido"),
        ("*Anatomia em Cadáveres* – disponível", "*Anatomia em Cadáveres*: disponível"),
    ]),
    "15febKjDEEupsWFSZXMB-fzTsVXtHAKzW": ("WhatsApp Lembrete", DOCX, [
        # o nome do vídeo é link (fica fora dos runs do parágrafo): o aviso entra antes dele
        ("Enviar vídeo:", "Enviar vídeo [REGRAVAR: o do 08/26 fala terça-feira e 12 mil inscritos]:"),
        ("Já passamos de 12 mil inscritos na Masterclass",
         "Já passamos de [NÚMERO REAL DE INSCRITOS EM 08/11] mil inscritos na Masterclass"),
    ]),
    "1-yQzs6JmU8TytN7tfao8LlI7MGdjdf-x": ("API Lembrete", DOCX, [
        ("Hoje, às 20h — Ao vivo no Zoom.", "Hoje, às 20h, ao vivo no Zoom."),
    ]),
}
PRA = re.compile(r"\bpra\b")


def trocar(p, velho, novo):
    """Troca um trecho que pode atravessar vários runs; o texto novo herda o formato do primeiro."""
    texto = "".join(r.text for r in p.runs)
    i = texto.find(velho)
    if i < 0: return False
    fim, pos = i + len(velho), 0
    primeiro = True
    for r in p.runs:
        a, b = pos, pos + len(r.text); pos = b
        if b <= i or a >= fim: continue
        ini_local, fim_local = max(i, a) - a, min(fim, b) - a
        r.text = r.text[:ini_local] + (novo if primeiro else "") + r.text[fim_local:]
        primeiro = False
    return True


def paragrafos(d):
    yield from d.paragraphs
    for t in d.tables:
        for linha in t.rows:
            for c in linha.cells: yield from c.paragraphs


if __name__ == "__main__":
    aplicar = "--aplicar" in sys.argv
    tk, _ = token()
    for fid, (nome, mime, trocas) in DOCS.items():
        url = f"{API}/{fid}/export?mimeType={DOCX}" if mime == GDOC else f"{API}/{fid}?alt=media"
        d = docx.Document(io.BytesIO(chamar(tk, "GET", url).read()))
        feitas, faltou = [], [v for v, _ in trocas]
        for p in paragrafos(d):
            for velho, novo in trocas:
                if trocar(p, velho, novo):
                    feitas.append(velho)
                    if velho in faltou: faltou.remove(velho)
            for r in p.runs:
                if PRA.search(r.text): r.text = PRA.sub("para", r.text); feitas.append("pra -> para")
        print(f"## {nome}: {len(feitas)} trocas" + (f" | NÃO ACHADO: {faltou}" if faltou else ""))
        if aplicar and feitas:
            buf = io.BytesIO(); d.save(buf)
            chamar(tk, "PATCH", f"{UPLOAD}/{fid}?uploadType=media&supportsAllDrives=true",
                   bruto=buf.getvalue(), cabecalhos={"Content-Type": DOCX})
            print("   SUBIU")
