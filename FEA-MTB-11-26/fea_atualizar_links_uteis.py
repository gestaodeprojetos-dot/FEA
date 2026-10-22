#!/usr/bin/env python3
"""Atualiza a planilha 'Links úteis' copiada do MTB 08/26 para o MTB 11/26.

Link que apontava para arquivo/pasta do Drive do 08/26 passa a apontar para a cópia nova (pelo id
de origem ou pelo nome). Link compartilhado entre eventos (páginas, UTMs, checkout) fica igual.
O que precisa ser criado de novo (grupo, tag, planilha de leads, Zoom) fica marcado na coluna C.
"""
import io, json, re, sys, unicodedata, urllib.parse
import openpyxl
from openpyxl.styles import Font, PatternFill
sys.path.insert(0, "/home/user/FEA/FEA-edicao-videos")
from fea_drive import token, chamar, API
from fea_duplicar_lancamento import novo_nome

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
UPLOAD = "https://www.googleapis.com/upload/drive/v3/files"
CLICKUP = "https://app.clickup.com/9013080622/v/f/1400270000003387/90130517360"
nfc = lambda s: unicodedata.normalize("NFC", s or "").strip()

mapa = json.load(open("FEA-mapa-drive-mtb-1126.json"))
por_nome = {}
for v in mapa.values():
    if not v["mime"].endswith("shortcut"):  # atalho aponta para arquivo compartilhado: mantém o link original
        por_nome.setdefault(nfc(v["nome"]), v)
raiz = mapa["1vvKGPIkXWWIoZHT1boqXFcZ6BcsxT9T1"]
PLANILHA = next(v for v in mapa.values() if nfc(v["nome"]).startswith(nfc("Links úteis")))["id"]


def url(v):
    m = v["mime"]
    if m.endswith("folder"): return f"https://drive.google.com/drive/folders/{v['id']}"
    if m.endswith("document"): return f"https://docs.google.com/document/d/{v['id']}/edit"
    if m.endswith("spreadsheet"): return f"https://docs.google.com/spreadsheets/d/{v['id']}/edit"
    return f"https://drive.google.com/file/d/{v['id']}/view"


def id_do_link(u):
    m = re.search(r"/(?:d|folders)/([A-Za-z0-9_-]{20,})", u or "")
    return m.group(1) if m else None


AVISO = PatternFill("solid", fgColor="FFF4CC")
GERAL = {  # linha: (texto, link, observação)
    2: ("11. DOWNSELL MTB 08/26", "https://drive.google.com/drive/folders/1vvKGPIkXWWIoZHT1boqXFcZ6BcsxT9T1", None),
    3: ("Links úteis . MTB - FEB 08/2026.xlsx", "https://drive.google.com/file/d/1GcdgnxIfnYk8-nz2QRZn49uhnRcuivGE/view", None),
    5: (raiz["nome"], url(raiz), None),
    7: ("ClickUp FEA-MTB 11/26", CLICKUP, None),
    14: ("A CRIAR: grupo SendFlow MTB 1126", None, "Júlio Couto (IAutomate). No 08/26: sndflw.com/i/masterclasstoxinabutolinica-0826"),
    16: ("A CRIAR: [1126] [MTB] - Lista de Inscrito", None, "Tag nova no ActiveCampaign (conta com o Mariano)"),
    18: ("A CRIAR: [JOÃO PITHON] MTB 1126 - CONTROLE DE LEADS", None, "A planilha do 08/26 não foi copiada (dado pessoal de lead)"),
    19: ("A CRIAR: [JOÃO PITHON] DADOS DE PESQUISA DE LEADS - MTB 1126", None, "Respostas da pesquisa do 08/26 não foram copiadas"),
    20: (None, None, "Conferir: o webhook do 08/26 aponta para 'febio-latam-pesquisa'"),
    23: ("A CONFIRMAR: link do Zoom de 09/11 às 20h", None, "Grazi cria o link. No 08/26: us06web.zoom.us/j/89745774248"),
    24: (None, None, "Confirmar oferta FEB em paralelo à Black Friday"),
    25: (None, None, "Confirmar oferta FEB em paralelo à Black Friday"),
}

if __name__ == "__main__":
    tk, _ = token()
    wb = openpyxl.load_workbook(io.BytesIO(chamar(tk, "GET", f"{API}/{PLANILHA}?alt=media").read()))
    log = []
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if ws.title == "GERAL" and c.column == 2 and c.row in GERAL: continue
                fid = id_do_link(c.hyperlink.target if c.hyperlink else None)
                alvo = mapa.get(fid) if fid else None
                if alvo and alvo["mime"].endswith("shortcut"): alvo = None
                if not alvo and c.hyperlink and isinstance(c.value, str):
                    alvo = por_nome.get(nfc(novo_nome(c.value)))
                if alvo and alvo is not raiz or (alvo is raiz and fid):
                    log.append(f"{ws.title}!{c.coordinate}: {c.value!r} -> {alvo['nome']!r}")
                    c.value, c.hyperlink = alvo["nome"], url(alvo)
                elif isinstance(c.value, str) and novo_nome(c.value) != c.value and not c.hyperlink:
                    log.append(f"{ws.title}!{c.coordinate}: {c.value!r} -> {novo_nome(c.value)!r}")
                    c.value = novo_nome(c.value)
    ws = wb["GERAL"]
    ws["C1"] = "Status MTB 11/26"; ws["C1"].font = Font(bold=True)
    for linha, (texto, link, obs) in GERAL.items():
        b = ws.cell(linha, 2)
        if texto is not None:
            b.value, b.hyperlink = texto, link
        if obs:
            ws.cell(linha, 3).value = obs
            ws.cell(linha, 3).fill = AVISO
            if texto and texto.startswith("A C"): b.fill = AVISO
        log.append(f"GERAL!B{linha}: {b.value!r} | {obs or ''}")
    ws.column_dimensions["C"].width = 70
    buf = io.BytesIO(); wb.save(buf)
    if "--aplicar" in sys.argv:
        chamar(tk, "PATCH", f"{UPLOAD}/{PLANILHA}?uploadType=media&supportsAllDrives=true", bruto=buf.getvalue(),
               cabecalhos={"Content-Type": XLSX})
        log.append("SUBIU")
    print("\n".join(log))
