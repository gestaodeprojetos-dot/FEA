"""Reescreve as legendas de vendas (03) e de remarketing de vendas (05) do MTB 11/26 para a Black Friday Vitalícia,
no mesmo formato dos docs do time. Fatos conferidos na página lp.joaopithon.com.br/black-friday/vitalicia/ em 28/10/2026.
Substitui o conteúdo dos mesmos Google Docs (mesmo link; histórico de versões do Docs guarda o anterior).
Uso: python3 fea_legendas_venda_black.py [--aplicar]"""
import io, sys
sys.path.insert(0, "/home/user/FEA/FEA-edicao-videos")
from fea_drive import token, chamar, API, UPLOAD
import docx
from docx.shared import Pt

SEP = "_" * 82
PRECO = "Lote Especial: 12x de R$ 568,52 (não alunos) ou 12x de R$ 465,09 (alunos). [conferir o lote vigente antes de publicar]"
BONUS = ["🎁 Pós-Graduação em Técnicas Injetáveis, com certificação",
         "🎁 Curso Presencial Full Face em 2027",
         "🎁 Ingresso para o Elite Injectors Congress 2027",
         "🎁 6 meses de Belasis, gestão completa da clínica",
         "🎁 Kit de Implementação, com entrega imediata",
         "🎁 Reology Map, 12 meses de acesso ao aplicativo de reologia",
         "🎁 Treinamento Padrão Ouro de Vendas"]

VENDAS = ("COPY DA LEGENDA DOS CRIATIVOS VENDAS I BLACK FRIDAY VITALÍCIA", "Versão", [
    ["VOCÊ VIU A AULA",
     "Na Masterclass em Toxina Botulínica, você viu como dose, ponto e profundidade mudam a previsibilidade do resultado.",
     "Na Black Friday Vitalícia, você leva a Formação Especialista em Botox e todos os meus outros cursos online, com acesso vitalício, inclusive aos cursos que eu ainda vou lançar.",
     PRECO,
     "Toque em \"Saiba Mais\" e garanta seu acesso vitalício."],
    ["ACESSO VITALÍCIO",
     "Preenchimento, botox, bioestimuladores, anatomia, fios, complicações e as imersões avançadas, em um único acesso vitalício.",
     "Você nunca mais vai precisar comprar ou renovar nenhum curso online meu.",
     "Toque em \"Saiba Mais\" e conheça a Black Friday Vitalícia."],
    ["7 BÔNUS INCLUSOS",
     "Além do acesso vitalício a todos os meus cursos online, a Black Friday Vitalícia inclui:",
     *BONUS,
     "Toque em \"Saiba Mais\" e veja a oferta completa."],
    ["GARANTIA DE 7 DIAS",
     "Entre, conheça os cursos por dentro e decida com segurança: a Black Friday Vitalícia tem garantia de 7 dias.",
     "Acesso vitalício a todos os meus cursos online atuais e futuros, por um único valor.",
     "Toque em \"Saiba Mais\" e garanta seu acesso vitalício."],
])
RMKT = ("COPY DAS LEGENDA DOS CRIATIVOS DE REMKT DE VENDAS I BLACK FRIDAY VITALÍCIA", "VERSÃO", [
    ["🔥 Acesso vitalício a todos os meus cursos online, atuais e futuros, em um único valor.",
     "Ao entrar na Black Friday Vitalícia, você ainda recebe:",
     *BONUS,
     "🔥 " + PRECO,
     "👉🏾 Toque em \"Saiba Mais\" e garanta seu acesso vitalício."],
    ["Ainda com dúvida sobre a Black Friday Vitalícia?",
     "Fale com o meu time pelo WhatsApp e entenda a condição para alunos e não alunos, as formas de pagamento e tudo o que está incluso.",
     "A oferta tem garantia de 7 dias.",
     "👉🏾 Toque em \"Saiba Mais\" e fale com o meu time."],
    ["⏳ A Black Friday Vitalícia vai até 20 de novembro. [confirmar dia e horário de fechamento do carrinho]",
     "Depois disso, o acesso vitalício a todos os meus cursos online não tem previsão de voltar.",
     "As formações em botox, preenchimento, bioestimuladores e anatomia e todos os outros cursos, mais 7 bônus, entre eles a Pós-Graduação em Técnicas Injetáveis.",
     "🔥 " + PRECO,
     "👉🏾 Toque em \"Saiba Mais\" e garanta seu acesso vitalício."],
])
DOCS = {"1kSwNYZJJeNkwDGQu4RcKboA-t51wAdi797LsohfK9Jc": VENDAS, "1NIhE98bnZw4yjS39B6mnjJDVauk9gyo7KdJ0CuqS_sk": RMKT}


def montar(titulo, rotulo, versoes):
    d = docx.Document()
    st = d.styles["Normal"]; st.font.name = "Montserrat"; st.font.size = Pt(11)
    d.add_paragraph().add_run("JOÃO PITHON                                 MASTERCLASS EM TOXINA BOTULÍNICA").bold = True
    d.add_paragraph(SEP)
    d.add_paragraph().add_run(titulo).bold = True
    for i, v in enumerate(versoes, 1):
        d.add_paragraph().add_run(f"{rotulo} {i}").bold = True
        primeira = rotulo == "Versão"                       # legendas de vendas abrem com uma chamada em caixa alta
        for k, linha in enumerate(v):
            p = d.add_paragraph()
            r = p.add_run(linha); r.bold = primeira and k == 0
        d.add_paragraph(SEP)
    b = io.BytesIO(); d.save(b)
    texto = "\n".join(p.text for p in d.paragraphs)
    assert " pra " not in texto and "—" not in texto
    return b.getvalue(), texto


if __name__ == "__main__":
    tk, _ = token() if "--aplicar" in sys.argv else (None, None)
    for fid, (tit, rot, vers) in DOCS.items():
        dados, texto = montar(tit, rot, vers)
        open(f"/tmp/claude-0/-home-user-FEA/ce2ef478-67a4-5284-b1d6-9d4c8f98649b/scratchpad/leg/{fid}.txt", "w").write(texto)
        print("ok", tit)
        if tk:
            chamar(tk, "PATCH", f"{UPLOAD}/{fid}?uploadType=media&supportsAllDrives=true", bruto=dados,
                   cabecalhos={"Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"})
            novo = chamar(tk, "GET", f"{API}/{fid}/export?mimeType=text/plain").read().decode("utf-8-sig")
            print("   conferido no Drive:", "BLACK FRIDAY VITALÍCIA" in novo and "FEB" not in novo.replace("Formação Especialista em Botox", ""))
