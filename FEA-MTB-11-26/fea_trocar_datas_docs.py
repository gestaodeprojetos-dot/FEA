#!/usr/bin/env python3
"""Troca as datas das copys do MTB 08/26 para o MTB 11/26 dentro dos docs copiados no Drive.

Baixa cada doc (Google Doc exportado como .docx ou .docx nativo), troca o texto preservando a
formatação dos trechos, e sobe de volta no mesmo arquivo. Não mexe em link (URL fica intacta).
Uso: python3 fea_trocar_datas_docs.py mapa.json [--aplicar]
"""
import io, json, re, sys, urllib.parse
import docx
sys.path.insert(0, "/home/user/FEA/FEA-edicao-videos")
from fea_drive import token, chamar, API

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
GDOC = "application/vnd.google-apps.document"
UPLOAD = "https://www.googleapis.com/upload/drive/v3/files"

# dia de agosto -> (dia da semana em novembro, dia). Âncora no dia do evento (D): 11/08 -> 09/11.
DIAS = {"03": ("segunda", "02"), "04": ("segunda", "02"), "05": ("terça", "03"), "06": ("quarta", "04"),
        "07": ("quinta", "05"), "08": ("sexta", "06"), "09": ("sábado", "07"), "10": ("domingo", "08"),
        "11": ("segunda", "09"), "12": ("terça", "10"), "13": ("quarta", "11")}
SEMANA = r"(segunda|terça|terca|quarta|quinta|sexta|sábado|sabado|domingo)"


def caixa(modelo, novo):
    if modelo.isupper(): return novo.upper()
    if modelo[:1].isupper(): return novo[:1].upper() + novo[1:]
    return novo


def dia_semana(m):
    velho, feira, num = m.group(1), m.group(2), m.group(3)
    if num not in DIAS: return m.group(0)
    sem, dia = DIAS[num]
    if feira and sem not in ("sábado", "domingo"): sem += "-feira"
    return caixa(velho, sem) + m.group(0)[len(velho) + len(feira or ""):].replace(num, dia, 1)


# ATENÇÃO: não é idempotente (dia da semana troca de novo a cada rodada). Rodar só uma vez por doc.
REGRAS = [
    (re.compile(SEMANA + r"(-feira)?,\s*(\d{2})\b", re.I), dia_semana),
    (re.compile(r"\b11 (de) (agosto)\b", re.I), lambda m: "9 " + m.group(1) + " " + caixa(m.group(2), "novembro")),
    (re.compile(r"(próxima|proxima) segunda-feira, dia 9", re.I), None),  # só marcador, já correto
    (re.compile(r"\b03/08 A 11/08\b", re.I), lambda m: "02/11 A 09/11"),
    (re.compile(r"\b11/08\b"), lambda m: "09/11"), (re.compile(r"\b11\.08\b"), lambda m: "09.11"), (re.compile(r"\b10/08\b"), lambda m: "08/11"),
    (re.compile(r"\b07/08\b"), lambda m: "05/11"), (re.compile(r"\b03/08\b"), lambda m: "02/11"),
    (re.compile(r"\b12/08\b"), lambda m: "10/11"), (re.compile(r"\b13/08\b"), lambda m: "11/11"),
    (re.compile(r"\b(0[16]|10)/07/2026\b"), lambda m: "22/10/2026"),
    (re.compile(r"(até|ate) (quinta)(-feira)", re.I), lambda m: m.group(1) + " " + caixa(m.group(2), "quarta") + m.group(3)),
    (re.compile(r"08/2026"), lambda m: "11/2026"), (re.compile(r"\b08/26\b"), lambda m: "11/26"),
    (re.compile(r"\b07/26\b"), lambda m: "11/26"), (re.compile(r"\b0826\b"), lambda m: "1126"),
    (re.compile(r"\bAGO/26\b"), lambda m: "NOV/26"),
]
REGRAS = [(r, f) for r, f in REGRAS if f]
# regras que nunca reconvertem o próprio resultado: seguras para a varredura extra trecho a trecho
SIMPLES = REGRAS[1:]
W_T = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"


def varredura_extra(d, log):
    """Texto em caixa de texto, controle de conteúdo ou tabela aninhada, fora de d.paragraphs."""
    partes = [d.element.body] + [x for s in d.sections for x in (s.header._element, s.footer._element)]
    for parte in partes:
        for t in parte.iter(W_T):
            if not t.text: continue
            texto, achados = t.text, []
            for regra, func in SIMPLES:
                for m in regra.finditer(texto):
                    novo = func(m)
                    if novo == m.group(0) or em_url(texto, m.start()): continue
                    if any(m.start() < f and a < m.end() for a, f, _ in achados): continue
                    achados.append((m.start(), m.end(), novo)); log.append(f"{m.group(0)!r} -> {novo!r}")
            for a, f, novo in sorted(achados, reverse=True):
                texto = texto[:a] + novo + texto[f:]
            t.text = texto


def em_url(texto, ini):
    a = texto.rfind(" ", 0, ini) + 1
    b = texto.find(" ", ini); b = len(texto) if b < 0 else b
    tok = texto[a:b]
    return "://" in tok or "www." in tok or ".com" in tok


def trocar_paragrafo(p, log):
    runs = p.runs
    if not runs: return
    texto = "".join(r.text for r in runs)
    achados = []  # passada única no texto original, sem reconverter o que já foi trocado
    for regra, func in REGRAS:
        for m in regra.finditer(texto):
            novo = func(m)
            if novo == m.group(0) or em_url(texto, m.start()): continue
            if any(m.start() < f and a < m.end() for a, f, _, _ in achados): continue
            achados.append((m.start(), m.end(), m.group(0), novo))
    for ini, fim, velho, novo in sorted(achados, reverse=True):
        log.append(f"{velho!r} -> {novo!r}")
        pos, colocado = 0, False
        for r in runs:
            a, b = pos, pos + len(r.text); pos = b
            if b <= ini or a >= fim: continue
            s, e = max(ini, a) - a, min(fim, b) - a
            r.text = r.text[:s] + ("" if colocado else novo) + r.text[e:]
            colocado = True


def paragrafos(d):
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs
    for s in d.sections:
        for parte in (s.header, s.footer):
            yield from parte.paragraphs


if __name__ == "__main__":
    mapa, aplicar = json.load(open(sys.argv[1])), "--aplicar" in sys.argv
    tk, _ = token()
    for v in mapa.values():
        if v["mime"] not in (DOCX, GDOC): continue
        url = (f"{API}/{v['id']}/export?mimeType=" + urllib.parse.quote(DOCX)) if v["mime"] == GDOC else f"{API}/{v['id']}?alt=media"
        d = docx.Document(io.BytesIO(chamar(tk, "GET", url).read()))
        log, vistos = [], set()
        for p in paragrafos(d):
            if p._p in vistos: continue
            vistos.add(p._p); trocar_paragrafo(p, log)
        varredura_extra(d, log)
        if not log: continue
        print(f"\n## {v['caminho']} ({len(log)} trocas)")
        for l in sorted(set(log)): print("   ", l, f"x{log.count(l)}")
        if aplicar:
            buf = io.BytesIO(); d.save(buf)
            chamar(tk, "PATCH", f"{UPLOAD}/{v['id']}?uploadType=media&supportsAllDrives=true",
                   bruto=buf.getvalue(), cabecalhos={"Content-Type": DOCX})
            print("    SUBIU")
