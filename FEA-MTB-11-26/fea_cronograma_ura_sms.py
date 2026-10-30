"""Cronograma de URA + SMS do MTB 11/26 (doc 7. URA + SMS/URA + SMS.docx), montado a partir do MTB 08/26
e das regras do lançamento pago mais recente (IFF +10ml 09/26). Regrava o mesmo arquivo (mesmo link),
fixando a revisão anterior para rollback. Uso: python3 fea_cronograma_ura_sms.py [--aplicar]"""
import io, json, sys
sys.path.insert(0, "/home/user/FEA/FEA-edicao-videos")
from fea_drive import token, chamar, API, UPLOAD
import docx
from docx.shared import Pt, RGBColor
from docx.oxml.shared import OxmlElement, qn

DOC = "1FlC3LQ40jvZ796gSWUW5dzol8Yer5pbh"
AUDIO = "https://drive.google.com/file/d/1Qs2M72axwQaA6Ezl-L1GjG9439Nc_4lb/view?usp=drive_link"
REGRAS = ["Enviar um SMS teste antes de cada disparo, para verificação.",
          "URA: áudio de no máximo 30 segundos, em MP3, sem cortes.",
          "SMS: no máximo 160 caracteres, incluindo link.",
          "Lista de telefones: exportar os inscritos do MTB 11/26 no dia do disparo (a planilha da pasta ainda é a lista do 08/26)."]
DIAS = [
    ("DATA: 08/11 (Domingo)", [
        ("SMS - 09h", "SMS Convencional", "09h00",
         "João Pithon: É amanhã! Masterclass em Toxina Botulínica, dia 9 de novembro, às 20h, ao vivo e sem replay no Zoom. Programe-se.", None, True)]),
    ("DATA: 09/11 (Segunda)", [
        ("SMS - 12h", "SMS", "12h00",
         "João Pithon: COMEÇA HOJE! Masterclass em Toxina Botulínica. Aprenda 9 técnicas avançadas de toxina, ao vivo e sem replay, às 20h no Zoom!", None, False),
        ("SMS - 19h", "SMS", "19h00",
         "João Pithon: FALTA 1 HORA! A Masterclass em Toxina Botulínica começa às 20h no Zoom. O link está no grupo do WhatsApp e no seu e-mail.", None, True),
        ("URA - 20h", "URA", "20h00", None, "AO-VIVO-MTB.mp3", False)]),
]


def link(par, texto, url):
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "1155CC"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    t = OxmlElement("w:t"); t.text = texto; r.append(rpr); r.append(t); h.append(r); par._p.append(h)


def montar():
    d = docx.Document()
    st = d.styles["Normal"]; st.font.name = "Montserrat"; st.font.size = Pt(11)
    p = d.add_paragraph(); r = p.add_run("MTB . 9 de novembro, às 20h"); r.bold = True; r.font.size = Pt(16)
    p = d.add_paragraph(); r = p.add_run("Disparo de URA . SMS"); r.bold = True
    for regra in REGRAS: d.add_paragraph(regra, style="List Bullet")
    d.add_paragraph("Novo neste lançamento (padrão IFF +10ml 09/26): SMS na véspera e SMS de 'falta 1 hora'.").runs[0].italic = True
    for dia, envios in DIAS:
        p = d.add_paragraph(); r = p.add_run(dia); r.bold = True; r.font.color.rgb = RGBColor(0x5B, 0x45, 0x99)
        for titulo, formato, hora, texto, audio, novo in envios:
            p = d.add_paragraph(); p.add_run(titulo + (" (novo)" if novo else "")).bold = True
            d.add_paragraph(f"Formato: {formato}")
            d.add_paragraph(f"Hora: {hora}")
            if texto:
                assert len(texto) <= 160, (titulo, len(texto))
                d.add_paragraph(texto); d.add_paragraph(f"({len(texto)} caracteres)").runs[0].italic = True
            if audio:
                link(d.add_paragraph("Áudio: "), audio, AUDIO)
    b = io.BytesIO(); d.save(b); return b.getvalue()


if __name__ == "__main__":
    dados = montar()
    open("/tmp/claude-0/-home-user-FEA/ce2ef478-67a4-5284-b1d6-9d4c8f98649b/scratchpad/ura2/cronograma.docx", "wb").write(dados)
    if "--aplicar" in sys.argv:
        tk, _ = token()
        m = json.load(chamar(tk, "GET", f"{API}/{DOC}?fields=headRevisionId&supportsAllDrives=true"))
        chamar(tk, "PATCH", f"{API}/{DOC}/revisions/{m['headRevisionId']}?supportsAllDrives=true", corpo={"keepForever": True})
        chamar(tk, "PATCH", f"{UPLOAD}/{DOC}?uploadType=media&supportsAllDrives=true", bruto=dados,
               cabecalhos={"Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"})
        print("subiu; rollback", m["headRevisionId"])
