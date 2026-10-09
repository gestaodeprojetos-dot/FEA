#!/usr/bin/env python3
"""Planilha "Criativos para Tráfego" do Ebook Olheiras LATAM, espelhando a aba Template
do modelo Brasil (Status, Campanha, Meio, Deadline, Fase, Legenda, Link Criativo, Link CTA),
com uma linha por criativo real e as referências LATAM e Brasil.

    python3 fea_planilha_trafego_latam.py SAIDA.xlsx
Sem datas: Deadline fica vazio até a gestão definir o sprint.
"""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROXO, LAVANDA, CINZA, VERMELHO, AZUL = '5B4599', 'D0D1E9', 'E9E9E9', 'BB3838', '3846BB'
borda = Border(*(Side(style='thin', color=CINZA),) * 4)
D = 'https://drive.google.com/drive/folders/'
F = 'https://drive.google.com/file/d/'
G = 'https://docs.google.com/document/d/'

REF = {
    'pasta_estaticos': D + '1RibhoTeWsYU5wf4q__1aV_QGXGCia2n4',
    'pasta_videos': D + '1ZKGD-pQ0iA8FHqx6Ekc8HvIWYeqkfPS3',
    'legenda_es': G + '1UgSAaYLfayStlF23n3jkKLxqM3u3I4481OtwTw4qPGg/edit',
    'copy_estaticos': G + '159seUx0p5uXMYmNENALTJBuB7QV1pv3KKey2U_wDD4c/edit',
    'briefing_criativos': G + '1SxAJCxfIiz_uqRXV0BuXKzhYxv-cDeL8NyhQVhSYAlU/edit',
    'pasta_brasil': D + '1UoFid4S3Fdb2VsZq49k_TZGHpf9clCNF',
    'planilha_brasil': F + '1iWdHQoTP6K0S507-DO3SnxldWlWT5Szq/view',
    'cta_brasil': 'https://eb.joaopithon.com.br/ebook-pto-a/',
}

BR = {  # arquivo Brasil -> fileId
    'Ads 01 - PTO - Feed.png': '1uVt5lFYt35AB-PiQS48kPxWeolvOyvzq', 'Ads 01 - PTO - Story.png': '1JA420MGAz2fX8xzSOKARSsi-MZHP67-N',
    'Ads 02 - PTO - Feed.png': '1mpt_KOOFurryoEQrHbcUB8x68n8z-ZHJ', 'Ads 02 - PTO - Story.png': '15Q1vLPguqyYAVCvPChSgifUMIfsR62rA',
    'Ads 03 - PTO - Feed.png': '1Xd3_nv7BhCLm9j-ku_qEocSj5Jft7nwk', 'Ads 03 - PTO - Story.png': '12MABWKzvjT6aZu8yQQsVp2XFRVfvjRKF',
    'Ads 04 - PTO - Feed.png': '1l3bCMMUBLz3JPpx9bMDWfDIW5PazXU_M', 'Ads 04 - PTO - Story.png': '1UP7Je1g-edTudBtIs5zmW_8NcijT-9U-',
    'Ads 05 - PTO - Feed.png': '1BJmhWAmX2ZrNNKZonrCPY5MlJyWmtmZe', 'Ads 05 - PTO - Story.png': '1A-Bke14HHbVaOD-HA0EFLNdrbFeOP93y',
    'Ads 06 - PTO - Feed.png': '18t2vaC7XPvtPpRrIBKmPbAHVLTWoY0i-', 'Ads 06 - PTO - Story.png': '1b2auyFzJLDKqNtY0jCJ1EV7NP8LbUHkX',
    'Ads 07 - PTO - Feed.png': '1YhACYv-Euv7JpTPJGC_P0HldYrJHm2RV', 'Ads 07 - PTO - Story.png': '1sfIUcxMXZOBE6E4eaT1Zk_t0cpoGgs_a',
    'Ads 08 - PTO - Feed.png': '1U0Mv9Lxz5aZnDFRPIo5tlN5wgQVlKdyk', 'Ads 08 - PTO - Story.png': '18V6vlx9UR5DsBFQo7a74Z69yjzXVkV0k',
    'Ads 09 - PTO - Feed.png': '11OUZw8-ZW-U2gIl5Sawvf3XK9_5JYcgA', 'Ads 09 - PTO - Story.png': '1QYKoxifTFy9b8XRSmU8bVu4Wx9EWNTGv',
    'Ads 10 - PTO - Feed.png': '14nyIcNhZZGQH3VJAiE5CLjyIMeZT8BiC', 'Ads 10 - PTO - Story.png': '1tj-BmRYfXm1IwcgdphA9WaU2fdWRtR6E',
    'Ads 11 - PTO - Feed.png': '10A2ipUVLV7Y94EdnvnPefseDKM5yn_Vc', 'Ads 11 - PTO - Story.png': '1MtNBl-mAuoNiz1sQ06lsW6nF7nVF7Tr7',
    'Ads 12 - PTO - Feed.png': '104eSqgZgX6nV0wM8yXWuXHeEhXxK_I77', 'Ads 12 - PTO - Story.png': '1rwzSB0pY_d6XzCPf51LiYAEyxjtpM6K9',
    'Ads 13- PTO - Feed.png': '1zCandJ3jLzHvFvlaJe0LBslWleozP4Md', 'Ads 13 - PTO - Story.png': '1nn7w2Ts4JvHle6MwPLIDwdGHRN6iI7-c',
    'Ads 14 - PTO - Feed.png': '1QjdLkAzsugDWxDqEazt7J9iM6WI_Nrdp', 'Ads 14 - PTO - Story.png': '1fQQZTpvI1H0rfiztW65F0bZlxu8SJgbT',
    'Ads 15 - PTO - Feed.png': '1E0pQnLzpaNQkC8UxOAMHdzLPMsCWS0_p', 'Ads 15 - PTO - Story.png': '1v7c4O3oClVjvHN7wgYjh94BUhrVXdJYL',
    '[STORIES] ADS 01.jpg': '1NqfHjmTbmnHokPZol2sdZPXPTLFAQOZU', '[FEED] ADS 02.jpg': '15mabKbWCO3eKSvluV6Qm8uBlw6rHjYC0',
    '[FEED] ADS 03.jpg': '1XutwAdMMUIb6r1BHGpiWb-JeCPUErZjk', '[FEED] ADS 04.jpg': '1rLBELyMJpQagh0UEqivvqrOQBax3cLRB',
    '[STORIES] ADS 04.jpg': '1itU4M9hs1BCuq0HPHdPGal5fLHoVxOS3', '[FEED] ADS 05.jpg': '1ir6NT5joQzE4VJwvflRr7urmguT34HI7',
    '[STORIES] ADS 05.jpg': '1xTYe9qtkvS6ujw4Rbgb4MPr_6dP8c9mH', '[FEED] ADS 06.jpg': '14vwqVJEI7AvZZyL9bj58dNOlpSwPS4Ou',
    '[STORIES] ADS 06.jpg': '1EhGjZw43yQghAYa4WVuMDCyXNrh3nX6A', '[FEED] ADS 07.jpg': '1aKrMUsrNkKa9PX5ZjsksXwOcPmtbCUwE',
    '[STORIES] ADS 07.jpg': '1VvYlEPiprd5wMUJrlkRwHE-PF7rPaii8', '[FEED] ADS 08.jpg': '18LE9PfKT5oLwlqV3OA4lYkYzxU6wxvRf',
    '[STORIES] ADS 08.jpg': '16rACtcIVWMDw4tCrI9Dk3Dgi5gOD6VOK', '[FEED] ADS 09.jpg': '1mq0kkwI4zhri5MpiHP9j47lFiaLgGYGy',
    '[STORIES] ADS 09.jpg': '1seLOpLEPvt8h6xMT0m_4vRhG5xjdk_r-', '[FEED] ADS 10.jpg': '1FbdZ9y1ig4lZdiXca8koVJFCVkiqNA7F',
    '[STORIES] ADS 10.jpg': '1EJDhyb9OR8oMEVj-BEvVKmxeAcYcMq75', '[FEED] ADS 11.jpg': '14fLmZdyzqeWMjujtqHa2brLeheO9YCrJ',
    '[STORIES] ADS 11.jpg': '1k4Ir1JiONGuGxmJcMz4rVS_8fkDaPqYH', '[FEED] ADS 12.jpg': '13HDyklywbvy79y8wT-yLmWijzhvmtdIZ',
    '[STORIES] ADS 12.jpg': '168ABz0MOoVKTT1YPFLofXpB7pBCFf_mh', '[FEED] ADS 13.jpg': '19akeAE38SsbFFCK_apTlVlhrhOcPAISt',
    '[STORIES] ADS 13.jpg': '15Kgi-NCFcaAyLqrcp4RCGZlcJVFitPcm', '[FEED] ADS 14.jpg': '1lHsNgf1YVlVmstgy59dJQML12hexYUiM',
    '[FEED] ADS 15.jpg': '1NWjq-WryzfcuRP3C5j1hvxABU1h2YUhm', '[FEED] ADS 15_V2.jpg': '1K4PYPL8RDeO5yLB37bq31LxxdWD547Kr',
    '[FEED] ADS 16.jpg': '1i9f5ruBD1RX8UdG9iRgl0RmEb52oXSHw', '[FEED] ADS 17.jpg': '1tpZGYI-JB4gbZVV-ZE-tb-lGdNH8yH9e',
    '[FEED] ADS 18.jpg': '1AM4XYrsfrYx23aAq2kCdBtb9zynh97EZ', '[STORIES] ADS 18.jpg': '1tSWAbXUXROyXk0y_Tq5zSjllQmW41GMB',
    '[FEED] ADS 19.jpg': '1h_jcJLUtFwUiWBmzix9KBN5guvUZELH1', '[FEED] ADS 20.jpg': '1YsCyuPuXcHbQqpbSlUk0TqWd76bB8bhg',
    '[STORIES] ADS 20.jpg': '1n8AEtb2TqMC-kaAPU7kuMzjYSI7IEORZ',
    'Ads 31 Feed - PTO.png': '1Ip_aU25avakqgNqac2Nh0pgNGCB1DeTh', 'Ads 31 Story - PTO.png': '1GWpfgK8B40ac4JpY6AqpQq616BcGZzcW',
    'Ads 32 Feed - PTO.png': '16pLrAmmjXiWRsVwXVt5JMAEqId5kJw_o', 'Ads 32 Story - PTO.png': '1AzJnVcPuZQ0F1xWUfUZ0rGHajcXVPc8n',
    'Ads 33 Feed - PTO.png': '1-Ho0GMtpp8l9ojDlDdxhbYybHQiIC3n4', 'Ads 33 Story - PTO.png': '1lLH-Sd2G3kYFzuY8hqA-yMtbLnC7QHQq',
    'Ads 34 Feed - PTO.png': '1lOharMaN3GiqbyhzmHQ6UbyJo2U-y6QK', 'Ads 34 Story - PTO.png': '1PQQadGv8eD4FsYfuBvAEnCuMxTTPmJGe',
    'Ads 35 Feed - PTO.png': '1ayLT4OhbrKigvVDKCrs2krht8t1Foey4', 'Ads 35 Story - PTO.png': '1PE8VYvt4p2C3Ep4HqA5OADz2x11FLZ_Y',
    'Ads 36 Feed - PTO.png': '19X0xvj5ut2nU_mCw5MTAlvnbINj0Dki0', 'Ads 36 Story - PTO.png': '1lvRrd8K-hPsFXfVZc6K5SpNJez9OkSWF',
    'Ads 37 Feed - PTO.png': '1EO1kbyRmEwKK7md4FABoMB3yB_KY3cpQ', 'Ads 37 Story - PTO.png': '11KGvSj_wSQm46LImuaJIGmMjTOIll8pu',
    'Ads 38 Feed - PTO.png': '1BWNtCq1PUXgqu7TOWB7UuD7jZplHfs0y', 'Ads 38 Story - PTO.png': '1j-kqPxa1yb3iHHwqHk0nS56lCu9089JT',
    'Ads 39 Feed - PTO.png': '1_ieeCS-ppevrbW2MZvab6EFzQxl01TAp', 'Ads 39 Story - PTO.png': '1pqqWg9PzMcYAc-bdt7jdfns3wFnPkK0u',
    'Ads 40 Feed - PTO.png': '1Y1W3do1YYXLyd7LTAkFza-HtcNEgehFp', 'Ads 40 Story - PTO.png': '18X6Xlcta920XCjPCEDfYGxIJj_Ck9VsZ',
}
VIDEOS = [('Ads 16', '11rpcytGMIqX5kG3QS7GeONU1lFwil7_l'), ('Ads 17', '1SoJEDv7hkUu2D4XGP5NdFVeADzyk8zIb'), ('Ads 18', '10mTDrlQQ1wDdChSp_sdWnXLRK1fc8Lze'),
          ('Ads 19', '1VY8P2RbQOkvlEfo148x7gjBwqBWpkUNh'), ('Ads 20', '1LFw6td1cQGJZpPvCrs-YrNDlG6UDwD-X'), ('Ads 21', '1AvSKHifec46BsWe1pwUZfrxAKrsAq80X'),
          ('Ads 22', '1bsk_ie-egJdlJPyH7wtLW7Y6yQeNOC0j'), ('Ads 23', '10sZUrX3waNbqfdL4u2Ca85r1SIUpX5oh'), ('Ads 24', '1UCMzPnKK29VDVmz3zkekHiEhKmKoE216'),
          ('Ads 25', '1cYjLF-AxMYf6F9dQth_W1NjmwPKP_EGD'), ('Ads 26', '1J2MupmT7oGbmFDChDm7GM9d-_vwRE7oE'), ('Ads 27', '1PtjQkSjQpxZWrZhuZF6vzByBnYLcW4ad'),
          ('Ads 28', '1iza-ggQFdRPRKmE-WFHDpgQTBMLPWNQb'), ('Ads 29', '1J5NrGtb2lveHNd1eWgS6cPtlWcR3FDlC'), ('Ads 30', '1FgPZ5w5Qo_q2h4WeEtUvVxtrF2UQhvey')]

MOCKUP = {'Ads 13', 'Ads 14', 'Ads 15', '[FEED] ADS 02', '[STORIES] ADS 01', '[FEED] ADS 09', '[STORIES] ADS 09', '[FEED] ADS 15_V2', '[FEED] ADS 17', 'Ads 34', 'Ads 36'}


def nome_latam(br):
    base, ext = os.path.splitext(br)
    if base.startswith('['):
        return f'FEA-{base} - LATAM{ext}'
    return 'FEA-' + base.replace('Ads 13- PTO', 'Ads 13 - PTO').replace('PTO', 'PTO-LATAM') + ext


def chave(br):
    b = os.path.splitext(br)[0]
    if b.startswith('['):
        return b
    return b.split(' - ')[0].split(' Feed')[0].split(' Story')[0].replace('Ads 13-', 'Ads 13').strip()


def formato(br):
    return 'Story' if ('Story' in br or 'STORIES' in br) else 'Feed'


def celula(c, valor, rotulo=None):
    if isinstance(valor, str) and valor.startswith('http'):
        c.value = f'=HYPERLINK("{valor}","{(rotulo or "Abrir").replace(chr(34), chr(39))}")'
        c.font = Font(color=AZUL, underline='single')
    else:
        c.value = valor
        if isinstance(valor, str) and valor.lower().startswith('pendente'):
            c.font = Font(color=VERMELHO, bold=True)


def cabecalho(ws, cols, larg):
    ws.append(cols)
    for i, c in enumerate(ws[ws.max_row], 1):
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=ROXO)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = larg[i - 1]
    ws.freeze_panes = 'A2'


def main(saida):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Template'
    cols = ['Status', 'Campanha', 'Meio', 'Deadline', 'Fase', 'Legenda', 'Link Criativo', 'Link CTA',
            'Nome do criativo (LATAM)', 'Tipo', 'Formato', 'Copy do criativo', 'Referência Brasil', 'UTM', 'Observações']
    cabecalho(ws, cols, [16, 26, 11, 11, 14, 18, 18, 22, 40, 10, 9, 18, 18, 20, 48])
    camp = 'Ebook Olheiras LATAM - Perpétuo'
    linhas = []
    for br in BR:
        k = chave(br)
        obs = 'Preço em dólar pendente (arte com placeholder US$ até a gestão definir).'
        status = 'Em produção'
        if k in MOCKUP:
            obs += ' Mockup do ebook: trocar capa/página pela edição ES (fase 2).'
        if k in ('[FEED] ADS 05', '[STORIES] ADS 05', '[FEED] ADS 08', '[STORIES] ADS 08'):
            obs += ' Depoimento em print fica em PT com legenda ES.'
        if k in ('Ads 06', 'Ads 09'):
            status = 'Em revisão'
        linhas.append([status, camp, 'Meta Ads', None, 'Vendas ebook', REF['legenda_es'], REF['pasta_estaticos'], 'pendente: página LATAM (André Vivas)',
                       nome_latam(br), 'Estático', formato(br), REF['copy_estaticos'], F + BR[br] + '/view', 'pendente (Matheus Saar)', obs])
    for nome, vid in VIDEOS:
        linhas.append(['Em tradução (HeyGen)', camp, 'Meta Ads', None, 'Vendas ebook', REF['legenda_es'], REF['pasta_videos'], 'pendente: página LATAM (André Vivas)',
                       f'FEA-{nome} - PTO-LATAM.mp4', 'Vídeo', 'Vertical', 'Transcrição ES na pasta VÍDEOS', F + vid + '/view', 'pendente (Matheus Saar)',
                       'Dublado no HeyGen (modo precision, espanhol LATAM). Conferir sincronia labial e legenda antes de subir.'])
    rotulos = {5: 'Leyenda ES', 6: 'Pasta LATAM', 11: 'Copy ES', 12: 'Original BR'}
    for l in linhas:
        ws.append([None] * len(cols))
        r = ws[ws.max_row]
        for i, v in enumerate(l):
            celula(r[i], v, rotulos.get(i))
            r[i].border = borda
            r[i].alignment = Alignment(vertical='center', wrap_text=i in (8, 14))
    dv = DataValidation(type='list', formula1='"A produzir,Em produção,Em revisão,Em tradução (HeyGen),Agendar,No ar,Pausado"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'A2:A{ws.max_row}')

    rf = wb.create_sheet('Referências')
    cabecalho(rf, ['Item', 'Link'], [44, 70])
    for nome, chave_ in [('Pasta ESTÁTICOS (LATAM)', 'pasta_estaticos'), ('Pasta VÍDEOS (LATAM)', 'pasta_videos'),
                         ('Legendas dos criativos (ES)', 'legenda_es'), ('Copy dos criativos estáticos PT | ES v2', 'copy_estaticos'),
                         ('Briefing de criativos e edição de vídeo LATAM', 'briefing_criativos'), ('Criativos Brasil (1. Vendas)', 'pasta_brasil'),
                         ('Planilha modelo Brasil', 'planilha_brasil'), ('Página de vendas Brasil (referência de CTA)', 'cta_brasil')]:
        rf.append([nome])
        celula(rf[rf.max_row][1], REF[chave_], 'Abrir')
    rf.append([])
    rf.append(['Página de vendas LATAM']); celula(rf[rf.max_row][1], 'pendente: André Vivas')
    rf.append(['Checkout Hotmart USD']); celula(rf[rf.max_row][1], 'pendente: gestão')
    rf.append(['Preço em dólar do ebook']); celula(rf[rf.max_row][1], 'pendente: gestão')

    cp = wb.create_sheet('Como preencher')
    cabecalho(cp, ['Status', 'Quando usar'], [24, 80])
    for s in [('A produzir', 'arte ainda não iniciada'), ('Em produção', 'arte em tradução ou com placeholder de preço'),
              ('Em revisão', 'arte pronta aguardando aprovação'), ('Em tradução (HeyGen)', 'vídeo na fila de dublagem'),
              ('Agendar', 'aprovado, falta subir no gerenciador'), ('No ar', 'anúncio ativo'), ('Pausado', 'desligado por performance ou fim de fase')]:
        cp.append(s)
    cp.append([])
    cp.append(['Regra', 'Deadline fica vazio até a gestão definir o sprint. Nome do criativo igual ao arquivo salvo na pasta LATAM.'])
    wb.save(saida)
    print('ok', saida, len(linhas), 'criativos')


if __name__ == '__main__':
    main(sys.argv[1])
